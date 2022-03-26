#!/usr/bin/env python
#!encoding:utf-8

from docx import Document
import time
import xlsxwriter
from glob import glob
import re
import os
import win32com.client as win32
from win32com.client import constants
from pathlib import Path
from typing import List

# Create list of paths to .doc files
# https://stackoverflow.com/questions/16056793/python-win32com-client-dispatch-looping-through-word-documents-and-export-to-pdf

current_file_path = Path(__file__).parent.absolute()
# output_dir = current_file_path / "output"

timestr = time.strftime("%Y%m%d-%H%M%S")

# convert doc to docx
def save_as_docx(filelist: List[Path], output_dir: Path) -> List[Path]:
    # Opening MS Word/wps
    word = win32.gencache.EnsureDispatch('kwps.Application')
    word.Visible = False
    res = []
    for file in filelist:
        # Rename path with .docx
        file_name = file.stem + '.docx'
        docx_path = output_dir / file_name
        res.append(docx_path)
        if(file.suffix != ".doc" or docx_path.exists()):
            continue
        doc = word.Documents.Open(str(file))
        doc.Activate()

        # Save and Close
        word.ActiveDocument.SaveAs(
            str(docx_path), FileFormat=constants.wdFormatXMLDocument
        )
        doc.Close(False)
    word.Quit()
    return res


def write_xlsx(to_write_data: List[List[str]], output_dir: Path, output_name: str = (timestr + ".xlsx")):
    # to_write = [[filepath.name, time_txt, "\r\n".join(content)]]
    xlsx_name = output_dir / output_name
    with xlsxwriter.Workbook(str(xlsx_name)) as workbook:
        worksheet = workbook.add_worksheet()
        worksheet.write_row(0, 0, ['文件名', '会议时间', '会议议题'])

        for row_num, data in enumerate(to_write_data):
            worksheet.write_row(row_num + 1, 0, data)


def extract_a_file(filepath: Path) -> List[str]:
    doc1 = Document(str(filepath))
    # print("dealing with {}".format(filepath))
    pl = [paragraph.text for paragraph in doc1.paragraphs]
    topic_begin = False
    content = []
    for line in pl:
        pass
    return [filepath.name, time_txt, "\r\n".join(content)]

if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument('--input_dir', metavar='dir', type=str, required=True,
                        help='input dir')
    parser.add_argument('--output_dir', metavar='dir', type=str, default="",
                        help='output dir')
    args = parser.parse_args()

    input_dir = Path(args.input_dir).absolute()
    if(args.output_dir == ""):
        output_dir = input_dir / "output"
    else:
        output_dir = Path(args.output_dir).absolute()
    output_dir.mkdir(parents=True, exist_ok=True)

    doc_files = input_dir.glob("**/*.doc")
    docx_files = save_as_docx(doc_files, output_dir)
    # print(docx_files)
    # print(len(list(doc_files)))
    # print(len(docx_files))
    # assert len(docx_files) == len(list(doc_files)), "wrong file numbers"
    content_list = []
    from tqdm import tqdm
    for docx in tqdm(docx_files):
        content_list.append(extract_a_file(docx))
    write_xlsx(content_list, output_dir)
